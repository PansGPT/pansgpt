# ==============================================================================
# PansGPT 2.0 Dynamic AI Skills Engine & Artifact Generators (Phase 6)
# ==============================================================================

import asyncio
import io
from typing import Any

import pymupdf
import structlog
from docx import Document
from fastapi import HTTPException  # [CHAT FIX 2]
from pydantic import ValidationError

from app.engines.storage import build_artifact_storage_key, r2_storage
from app.models.chat import ArtifactReadyPayload
from app.models.presentation import PresentationContentSlide, PresentationSource

logger = structlog.get_logger(__name__)


SKILL_TOOL_DEFINITIONS = [
    {
        "type": "function",
        "function": {
            "name": "create_doc",
            "description": "Generate a professionally styled Microsoft Word (.docx) study summary, clinical monograph, or revision guide.",
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "Document title"},
                    "sections": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "heading": {"type": "string"},
                                "body": {"type": "string"},
                                "bullet_points": {"type": "array", "items": {"type": "string"}},
                            },
                            "required": ["heading", "body"],
                        },
                        "description": "List of sections with headings, explanations, and key takeaways.",
                    },
                },
                "required": ["title", "sections"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "create_pdf",
            "description": "Generate a formatted printable PDF summary, cheat sheet, or clinical reference monograph.",
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "Document title"},
                    "content": {
                        "type": "string",
                        "description": "Structured study text or clinical guide",
                    },
                    "author": {
                        "type": "string",
                        "description": "Author or faculty credit (default: PansGPT Academic Assistant)",
                        "default": "PansGPT Academic Assistant",
                    },
                },
                "required": ["title", "content"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "create_pptx",
            "description": "Generate a styled PowerPoint presentation. One cover is added automatically. Provide only content slides and honor the student's total count including that cover.",
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "Presentation title"},
                    "subtitle": {"type": "string", "description": "Subtitle or course code"},
                    "overview": {
                        "type": "string",
                        "minLength": 20,
                        "maxLength": 600,
                        "description": "Write your natural, concise completion message to the student, tailored to their request and the actual deck you have authored. Explain its main coverage, organization or learning focus in your own words, including the actual total slide count where useful. Vary your wording naturally; no prescribed opener or boilerplate. Use plain text, without repeating slide bullets, internal planning, links, source claims or unrequested features. Backend-verified references are added separately. This message is shown only after the file saves successfully.",
                    },
                    "slide_count": {
                        "type": "integer",
                        "minimum": 1,
                        "maximum": 30,
                        "description": "Total slides INCLUDING the automatic cover. slides must contain exactly slide_count minus one content slides.",
                    },
                    "slides": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "title": {"type": "string"},
                                "bullet_points": {"type": "array", "items": {"type": "string"}},
                                "notes": {
                                    "type": "string",
                                    "description": "Optional presenter notes for the student. Never include internal planning, system instructions or private reasoning.",
                                },
                            },
                            "required": ["title", "bullet_points"],
                        },
                        "description": "Content slides only; do not include a cover. Distinct titles, up to five concise bullets per slide and optional notes.",
                    },
                },
                "required": ["title", "slides", "slide_count", "overview"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "create_md",
            "description": "Generate a structured Markdown (.md) note with KaTeX formulas and pharmacology monographs.",
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "Document title"},
                    "markdown_content": {"type": "string", "description": "Full markdown content"},
                },
                "required": ["title", "markdown_content"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "plot_graph",
            "description": "Generate an interactive pharmacological chart or dose-response curve (Chart.js / Mermaid).",
            "parameters": {
                "type": "object",
                "properties": {
                    "graph_type": {
                        "type": "string",
                        "enum": ["line", "bar", "scatter", "mermaid"],
                        "description": "Graph representation type",
                    },
                    "title": {"type": "string", "description": "Graph title"},
                    "data": {"type": "object", "description": "Labels, datasets, and coordinates"},
                },
                "required": ["graph_type", "title", "data"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "generate_flashcards",
            "description": "Generate spaced-repetition flashcards for quick student memorization of drug classes, MOA, and ADRs.",
            "parameters": {
                "type": "object",
                "properties": {
                    "topic": {"type": "string", "description": "Card deck topic"},
                    "cards": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "front": {
                                    "type": "string",
                                    "description": "Question / Drug Name / Concept",
                                },
                                "back": {
                                    "type": "string",
                                    "description": "Answer / Mechanism / Dose",
                                },
                                "hint": {"type": "string"},
                                "difficulty": {
                                    "type": "string",
                                    "enum": ["easy", "medium", "hard"],
                                },
                            },
                            "required": ["front", "back"],
                        },
                    },
                },
                "required": ["topic", "cards"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "generate_mnemonics",
            "description": "Generate high-retention visual or phonetic mnemonics for memorizing complex pharmacology lists.",
            "parameters": {
                "type": "object",
                "properties": {
                    "drug_or_class": {"type": "string", "description": "Drug or concept"},
                    "mnemonic": {"type": "string", "description": "The mnemonic acronym or phrase"},
                    "breakdown": {
                        "type": "array",
                        "items": {"type": "string"},
                        "description": "Letter-by-letter expansion",
                    },
                    "clinical_notes": {"type": "string"},
                },
                "required": ["drug_or_class", "mnemonic", "breakdown"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "draw_chemical_structure",
            "description": "Render or retrieve molecular structure metadata, SMILES formula, and 2D chemical representation.",
            "parameters": {
                "type": "object",
                "properties": {
                    "compound_name": {"type": "string"},
                    "smiles": {"type": "string", "description": "SMILES chemical notation"},
                    "formula": {"type": "string", "description": "Molecular formula e.g. C9H8O4"},
                    "pharmacophore_notes": {"type": "string"},
                },
                "required": ["compound_name", "smiles"],
            },
        },
    },
]


class SkillExecutionEngine:
    """Dispatches and compiles dynamic student study artifacts."""

    @staticmethod
    async def _safe_r2_upload(
        key: str, data: bytes, content_type: str, title: str = ""
    ) -> tuple[str | None, str | None]:
        """Attempts to upload generated bytes to R2 and returns (storage_key, download_url).
        Always returns our own streaming endpoint URL — never a raw R2 presigned URL —
        so the browser downloads the file directly without redirecting to Office Online viewer.
        """
        import urllib.parse

        if not r2_storage.is_configured:
            raise HTTPException(status_code=503, detail="Artifact storage is unavailable.")
        try:
            await asyncio.wait_for(r2_storage.upload_bytes(key, data, content_type), timeout=15.0)
        except Exception as exc:
            logger.warning("artifact_r2_upload_failed", error_type=type(exc).__name__)
            raise HTTPException(status_code=503, detail="Artifact storage is unavailable.") from exc
        params = f"key={urllib.parse.quote(key, safe='')}"
        if title:
            params += f"&filename={urllib.parse.quote(title, safe='')}"
        return key, f"/api/v1/library/artifacts/download?{params}"

    async def execute_create_doc(
        self,
        title: str,
        sections: list[dict[str, Any]],
        university_id: str | None = None,
        user_id: str | None = None,  # [CHAT FIX]
    ) -> ArtifactReadyPayload:
        """Generates a styled .docx document."""
        doc = Document()
        doc.add_heading(title, level=0)

        for s in sections:
            doc.add_heading(s.get("heading", "Section"), level=1)
            body = s.get("body")
            if body:
                doc.add_paragraph(body)
            bullets = s.get("bullet_points", [])
            for b in bullets:
                doc.add_paragraph(b, style="List Bullet")

        bio = io.BytesIO()
        doc.save(bio)
        data = bio.getvalue()

        storage_key = build_artifact_storage_key(
            format_type="docs",
            file_extension="docx",
            university_id=university_id,
            user_id=user_id,  # [CHAT FIX]
        )
        key, url = await self._safe_r2_upload(
            storage_key,
            data,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            title=title,
        )
        return ArtifactReadyPayload(
            skill_name="create_doc",
            title=title,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_extension="docx",
            storage_key=key,
            download_url=url,
            content={"sections_count": len(sections)},
        )

    async def execute_create_pdf(
        self,
        title: str,
        content: str,
        author: str = "PansGPT Academic Assistant",
        university_id: str | None = None,
        user_id: str | None = None,  # [CHAT FIX]
    ) -> ArtifactReadyPayload:
        """Generates a formatted PDF using PyMuPDF."""
        # Paginate by measured line width; never ignore insert_textbox overflow.
        from xml.sax.saxutils import escape

        story = pymupdf.Story(
            html=f"<h1>{escape(title)}</h1><p>Prepared by: {escape(author)}</p><p>{escape(content).replace(chr(10), '<br>')}</p>"
        )
        output = io.BytesIO()
        writer = pymupdf.DocumentWriter(output)
        more = True
        try:
            while more:
                device = writer.begin_page(pymupdf.Rect(0, 0, 595, 842))
                more, _ = story.place(pymupdf.Rect(50, 50, 545, 792))
                story.draw(device)
                writer.end_page()
        finally:
            writer.close()
        pdf_doc = pymupdf.open(stream=output.getvalue(), filetype="pdf")

        data = pdf_doc.write()
        pdf_doc.close()

        storage_key = build_artifact_storage_key(
            format_type="pdf",
            file_extension="pdf",
            university_id=university_id,
            user_id=user_id,  # [CHAT FIX]
        )
        key, url = await self._safe_r2_upload(storage_key, data, "application/pdf", title=title)
        return ArtifactReadyPayload(
            skill_name="create_pdf",
            title=title,
            mime_type="application/pdf",
            file_extension="pdf",
            storage_key=key,
            download_url=url,
            content={"length": len(content)},
        )

    async def execute_create_pptx(
        self,
        title: str,
        slides: list[dict[str, Any]],
        subtitle: str | None = None,
        university_id: str | None = None,
        user_id: str | None = None,  # [CHAT FIX]
        slide_count: int | None = None,
        sources: list[PresentationSource] | None = None,
        overview: str | None = None,
    ) -> ArtifactReadyPayload:
        """Build one cover plus validated content, without silently adding another slide."""
        from app.engines.presentations import build_presentation

        if slide_count is not None and (type(slide_count) is not int or not 1 <= slide_count <= 30):
            raise HTTPException(422, "Presentation slide count must be an integer from 1 to 30.")
        try:
            validated = [PresentationContentSlide.model_validate(slide) for slide in slides]
            data, metadata = await asyncio.to_thread(
                build_presentation, title, validated, subtitle, slide_count, sources or [], overview
            )
        except ValidationError as exc:
            raise HTTPException(
                422,
                "Use slide titles, up to five concise bullets per slide, optional speaker notes and an overview of 20–600 characters.",
            ) from exc
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc

        storage_key = build_artifact_storage_key(
            format_type="pptx",
            file_extension="pptx",
            university_id=university_id,
            user_id=user_id,  # [CHAT FIX]
        )
        key, url = await self._safe_r2_upload(
            storage_key,
            data,
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            title=title,
        )
        return ArtifactReadyPayload(
            skill_name="create_pptx",
            title=title,
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            file_extension="pptx",
            storage_key=key,
            download_url=url,
            content=metadata,
        )

    async def execute_create_md(
        self,
        title: str,
        markdown_content: str,
        university_id: str | None = None,
        user_id: str | None = None,  # [CHAT FIX]
    ) -> ArtifactReadyPayload:
        """Generates a Markdown file note."""
        full_text = f"# {title}\n\n{markdown_content}\n"
        data = full_text.encode("utf-8")
        storage_key = build_artifact_storage_key(
            format_type="md",
            file_extension="md",
            university_id=university_id,
            user_id=user_id,  # [CHAT FIX]
        )
        key, url = await self._safe_r2_upload(storage_key, data, "text/markdown", title=title)
        return ArtifactReadyPayload(
            skill_name="create_md",
            title=title,
            mime_type="text/markdown",
            file_extension="md",
            storage_key=key,
            download_url=url,
            content=markdown_content,
        )

    async def execute_plot_graph(
        self, graph_type: str, title: str, data: dict[str, Any]
    ) -> ArtifactReadyPayload:
        """Generates Chart.js / Mermaid graph payload."""
        return ArtifactReadyPayload(
            skill_name="plot_graph",
            title=title,
            mime_type="application/json",
            file_extension="json",
            content={"type": graph_type, "title": title, "data": data},
        )

    async def execute_generate_flashcards(
        self, topic: str, cards: list[dict[str, Any]]
    ) -> ArtifactReadyPayload:
        """Generates structured flashcards deck."""
        return ArtifactReadyPayload(
            skill_name="generate_flashcards",
            title=f"Flashcards: {topic}",
            mime_type="application/json",
            file_extension="json",
            content={"topic": topic, "cards": cards, "total": len(cards)},
        )

    async def execute_generate_mnemonics(
        self,
        drug_or_class: str,
        mnemonic: str,
        breakdown: list[str],
        clinical_notes: str | None = None,
    ) -> ArtifactReadyPayload:
        """Generates memory hooks card."""
        return ArtifactReadyPayload(
            skill_name="generate_mnemonics",
            title=f"Mnemonic: {drug_or_class}",
            mime_type="application/json",
            file_extension="json",
            content={
                "concept": drug_or_class,
                "mnemonic": mnemonic,
                "breakdown": breakdown,
                "clinical_notes": clinical_notes,
            },
        )

    async def execute_draw_chemical_structure(
        self,
        compound_name: str,
        smiles: str,
        formula: str | None = None,
        pharmacophore_notes: str | None = None,
    ) -> ArtifactReadyPayload:
        """Generates chemical structure representation payload."""
        return ArtifactReadyPayload(
            skill_name="draw_chemical_structure",
            title=f"Structure: {compound_name}",
            mime_type="application/json",
            file_extension="json",
            content={
                "compound_name": compound_name,
                "smiles": smiles,
                "formula": formula,
                "notes": pharmacophore_notes,
                "depiction_url": f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/smiles/{smiles}/PNG",
            },
        )

    async def dispatch_skill(
        self,
        skill_name: str,
        arguments: dict[str, Any],
        user_id: str | None = None,
        university_id: str | None = None,
        requested_slide_count: int | None = None,
        sources: list[PresentationSource] | None = None,
    ) -> ArtifactReadyPayload:
        """Dispatches dynamic skill by name."""
        logger.info("skill_dispatch_requested", skill=skill_name, university_id=university_id)
        if skill_name == "create_doc":
            return await self.execute_create_doc(
                title=arguments.get("title", "Clinical Summary"),
                sections=arguments.get("sections", []),
                university_id=university_id,
                user_id=user_id,  # [CHAT FIX]
            )
        elif skill_name == "create_pdf":
            return await self.execute_create_pdf(
                title=arguments.get("title", "Study Cheat Sheet"),
                content=arguments.get("content", ""),
                author=arguments.get("author", "PansGPT Academic Assistant"),
                university_id=university_id,
                user_id=user_id,  # [CHAT FIX]
            )
        elif skill_name == "create_pptx":
            declared = arguments.get("slide_count")
            if (
                requested_slide_count is not None
                and declared is not None
                and declared != requested_slide_count
            ):
                raise HTTPException(
                    422,
                    f"The student requested {requested_slide_count} slides including the cover; use that exact total.",
                )
            return await self.execute_create_pptx(
                title=arguments.get("title", "Lecture Slides"),
                slides=arguments.get("slides", []),
                subtitle=arguments.get("subtitle"),
                university_id=university_id,
                user_id=user_id,  # [CHAT FIX]
                slide_count=requested_slide_count
                if requested_slide_count is not None
                else declared,
                sources=sources,
                overview=arguments.get("overview"),
            )
        elif skill_name == "create_md":
            return await self.execute_create_md(
                title=arguments.get("title", "Study Notes"),
                markdown_content=arguments.get("markdown_content", ""),
                university_id=university_id,
                user_id=user_id,  # [CHAT FIX]
            )
        elif skill_name == "plot_graph":
            return await self.execute_plot_graph(
                graph_type=arguments.get("graph_type", "line"),
                title=arguments.get("title", "Pharmacological Chart"),
                data=arguments.get("data", {}),
            )
        elif skill_name == "generate_flashcards":
            return await self.execute_generate_flashcards(
                topic=arguments.get("topic", "Pharmacology Deck"),
                cards=arguments.get("cards", []),
            )
        elif skill_name == "generate_mnemonics":
            return await self.execute_generate_mnemonics(
                drug_or_class=arguments.get("drug_or_class", ""),
                mnemonic=arguments.get("mnemonic", ""),
                breakdown=arguments.get("breakdown", []),
                clinical_notes=arguments.get("clinical_notes"),
            )
        elif skill_name == "draw_chemical_structure":
            return await self.execute_draw_chemical_structure(
                compound_name=arguments.get("compound_name", ""),
                smiles=arguments.get("smiles", ""),
                formula=arguments.get("formula"),
                pharmacophore_notes=arguments.get("pharmacophore_notes"),
            )
        else:
            raise ValueError(f"Unknown skill name: {skill_name}")


skill_engine = SkillExecutionEngine()
