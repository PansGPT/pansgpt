"""PowerPoint templates and bounded previews extracted from the owned saved file."""

import asyncio
import io
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path, PurePosixPath

import pymupdf
from lxml.etree import XMLSyntaxError
from pptx import Presentation
from pptx.dml.color import ColorFormat, RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt
from pydantic import TypeAdapter

from app.models.presentation import (
    PresentationContentSlide,
    PresentationOverview,
    PresentationPreviewResponse,
    PresentationShape,
    PresentationSlide,
    PresentationSource,
    PresentationText,
)

MAX_SOURCE_BYTES = 24 * 1024 * 1024
_slots = asyncio.Semaphore(2)


def _wrap(text: str, width: float, size: float, bold: bool = False) -> list[str]:
    font = pymupdf.Font("hebo" if bold else "helv")
    lines: list[str] = []
    for block in text.replace("\v", "\n").split("\n"):
        line = ""
        for word in block.split():
            candidate = f"{line} {word}".strip()
            if line and font.text_length(candidate, fontsize=size) > width:
                lines.append(line)
                line = ""
            for character in word:
                if font.text_length(line + character, fontsize=size) > width and line:
                    lines.append(line)
                    line = ""
                line += character
            line += " "
        lines.append(line.rstrip())
    return lines or [""]


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#"))


def build_presentation(
    title: str,
    slides: list[PresentationContentSlide],
    subtitle: str | None,
    slide_count: int | None,
    sources: list[PresentationSource],
    overview: str | None = None,
) -> tuple[bytes, dict[str, object]]:
    if overview is not None:
        overview = TypeAdapter(PresentationOverview).validate_python(
            overview.strip() if isinstance(overview, str) else overview
        )
    if not title.strip() or len(title) > 160:
        raise ValueError("Use a presentation title of 1–160 characters.")
    total = len(slides) + 1
    if not 1 <= total <= 30 or (slide_count is not None and slide_count != total):
        raise ValueError(
            f"The deck requires {slide_count or '1–30'} slides including one cover; "
            f"received {total}. Supply content slides only, without another cover."
        )
    titles = [" ".join(title.casefold().split())]
    for slide in slides:
        normalized = " ".join(slide.title.casefold().split())
        if not normalized or normalized in titles:
            raise ValueError("Give each content slide a distinct title; the cover is added once.")
        titles.append(normalized)
        if any(not bullet.strip() for bullet in slide.bullet_points):
            raise ValueError("Slide bullet points cannot be blank.")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = title
    prs.core_properties.author = "PansGPT"
    prs.core_properties.subject = "PansGPT study presentation"
    source_note = (
        "Sources\n" + "\n".join(f"• {s.title}" + (f" — {s.url}" if s.url else "") for s in sources)
        if sources
        else "Based on the conversation. No external sources were attached to this deck."
    )

    def box(slide: Slide, x: float, y: float, w: float, h: float, color: str) -> None:
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(color)
        shape.line.fill.background()

    def text(
        slide: Slide,
        value: str,
        x: float,
        y: float,
        w: float,
        h: float,
        size: float,
        color: str,
        bold: bool = False,
        name: str = "",
        align: PP_ALIGN = PP_ALIGN.LEFT,
    ) -> None:
        while size > 14 and len(_wrap(value, w * 72, size, bold)) * size * 1.22 > h * 72:
            size -= 1
        if len(_wrap(value, w * 72, size, bold)) * size * 1.22 > h * 72:
            raise ValueError("There is too much text on a slide. Shorten its bullet points.")
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            shape.name = name
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        paragraph.line_spacing = 1.22
        paragraph.space_after = Pt(0)
        paragraph.text = value
        for run in paragraph.runs:
            run.font.name, run.font.size, run.font.bold = "Arial", Pt(size), bold
            run.font.color.rgb = _rgb(color)

    for index in range(total):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        dark = index == 0 or (total > 3 and index == total - 1)
        background, foreground, muted = (
            ("#151515", "#FAFAF9", "#B7B7B7") if dark else ("#FAFAF9", "#171717", "#666666")
        )
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(background)
        text(slide, "PANSGPT  /  STUDY PRESENTATION", 0.85, 0.48, 9, 0.35, 10, "#12B76A", True)
        text(slide, "PansGPT", 0.85, 6.95, 5, 0.22, 9, muted)
        text(
            slide,
            f"{index + 1:02d} / {total:02d}",
            11.0,
            6.95,
            1.45,
            0.22,
            9,
            muted,
            align=PP_ALIGN.RIGHT,
        )
        if index == 0:
            box(slide, 0.85, 1.65, 0.75, 0.055, "#12B76A")
            text(slide, title, 0.85, 2.05, 9.25, 2.2, 46, foreground, True, "pansgpt-title")
            text(
                slide,
                subtitle or "A focused presentation for your studies.",
                0.9,
                4.65,
                9.2,
                0.9,
                21,
                muted,
            )
            text(
                slide,
                f"{total:02d}",
                10.7,
                2.2,
                1.6,
                1.2,
                60,
                "#12B76A",
                True,
                align=PP_ALIGN.CENTER,
            )
            text(slide, "SLIDES", 10.7, 3.5, 1.6, 0.4, 11, muted, align=PP_ALIGN.CENTER)
            slide.notes_slide.notes_text_frame.text = source_note
        else:
            content = slides[index - 1]
            text(slide, content.title, 0.85, 1.05, 11.5, 1.0, 32, foreground, True, "pansgpt-title")
            box(slide, 0.85, 2.17, 0.65, 0.045, "#12B76A")
            row_height = min(1.1, 4.1 / len(content.bullet_points))
            for number, bullet in enumerate(content.bullet_points):
                y = 2.48 + number * row_height
                box(slide, 0.89, y + 0.12, 0.065, 0.17, "#12B76A")
                text(slide, bullet, 1.2, y, 10.85, row_height - 0.13, 22, foreground)
            slide.notes_slide.notes_text_frame.text = (content.notes + "\n\n" + source_note).strip()
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue(), {
        "slides_count": total,
        "slide_titles": [title, *[slide.title for slide in slides]],
        "sources": [source.model_dump() for source in sources],
        "source_note": "Sources attached" if sources else "No external sources attached",
        "preview_available": True,
        "overview": overview,
    }


def _color(color: ColorFormat, fallback: str) -> str:
    try:
        if color.rgb is not None:
            return f"#{color.rgb}"
    except (AttributeError, TypeError):
        pass
    return fallback


def read_presentation(data: bytes) -> PresentationPreviewResponse:
    if len(data) > MAX_SOURCE_BYTES:
        raise ValueError("The presentation is too large to preview.")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        parts = archive.infolist()
        if len(parts) > 600 or sum(part.file_size for part in parts) > 64 * 1024 * 1024:
            raise ValueError("Presentation package limits exceeded.")
        for part in parts:
            if (
                part.flag_bits & 1
                or part.file_size > 16 * 1024 * 1024
                or ".." in PurePosixPath(part.filename).parts
                or "\\" in part.filename
            ):
                raise ValueError("Invalid presentation package.")
            if part.filename.endswith(".xml"):
                xml = archive.read(part).upper()
                if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                    raise ValueError("Unsafe presentation XML.")
    prs = Presentation(io.BytesIO(data))
    if not 1 <= len(prs.slides) <= 30:
        raise ValueError("Preview supports presentations of 1–30 slides.")
    warnings: set[str] = set()
    if prs.core_properties.subject != "PansGPT study presentation":
        warnings.add(
            "Some PowerPoint theme formatting may differ in this preview. Download for the original formatting."
        )
    pages: list[PresentationSlide] = []
    for index, slide in enumerate(prs.slides):
        if len(slide.shapes) > 150:
            raise ValueError("Too many elements on a slide.")
        elements: list[PresentationShape | PresentationText] = []
        title = slide.shapes.title.text if slide.shapes.title else f"Slide {index + 1}"
        for shape in slide.shapes:
            x, y, width, height = (
                float(value) / 12700 for value in (shape.left, shape.top, shape.width, shape.height)
            )
            if (
                shape.shape_type
                not in {
                    MSO_SHAPE_TYPE.AUTO_SHAPE,
                    MSO_SHAPE_TYPE.TEXT_BOX,
                    MSO_SHAPE_TYPE.PLACEHOLDER,
                }
                or shape.rotation
            ):
                warnings.add(
                    "Some media or rotated objects are not shown. Download the file to view all content."
                )
                continue
            fill = _color(shape.fill.fore_color, "#FFFFFF") if shape.fill.type == 1 else None
            stroke = _color(shape.line.color, "#888888") if shape.line.fill.type == 1 else None
            if fill or stroke:
                ellipse = (
                    shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                    and shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.OVAL
                )
                elements.append(
                    PresentationShape(
                        kind="ellipse" if ellipse else "rect",
                        x=x,
                        y=y,
                        width=width,
                        height=height,
                        fill=fill,
                        stroke=stroke,
                        stroke_width=float(shape.line.width or 12700) / 12700,
                        radius=0.0,
                    )
                )
            if not shape.has_text_frame:
                continue
            if shape.name == "pansgpt-title":
                title = shape.text
            frame = shape.text_frame
            left = x + float(frame.margin_left) / 12700
            available = max(1.0, width - float(frame.margin_left + frame.margin_right) / 12700)
            cursor = y + float(frame.margin_top) / 12700
            for paragraph in frame.paragraphs:
                if not paragraph.text:
                    continue
                if len(paragraph.text) > 6000:
                    raise ValueError("A slide has too much text to preview.")
                font = paragraph.runs[0].font if paragraph.runs else paragraph.font
                size = float((font.size or paragraph.font.size or Pt(20)).pt)
                bold, italic = (
                    bool(font.bold or paragraph.font.bold),
                    bool(font.italic or paragraph.font.italic),
                )
                color = _color(font.color, "#171717")
                align = {PP_ALIGN.CENTER: "middle", PP_ALIGN.RIGHT: "end"}.get(
                    paragraph.alignment, "start"
                )
                position = left + (
                    available / 2 if align == "middle" else available if align == "end" else 0
                )
                cursor += float(paragraph.space_before or 0) / 12700
                spacing = paragraph.line_spacing
                line_height = (
                    float(spacing.pt) if hasattr(spacing, "pt") else size * float(spacing or 1.22)
                )
                for line in _wrap(paragraph.text, available, size, bold):
                    if cursor + size > y + height + 2:
                        warnings.add(
                            "Some text exceeds its slide box. Download the file to inspect its layout."
                        )
                        break
                    elements.append(
                        PresentationText(
                            x=position,
                            y=cursor + size,
                            text=line,
                            font_size=size,
                            bold=bold,
                            italic=italic,
                            color=color,
                            align=align,
                        )
                    )
                    cursor += line_height
                cursor += float(paragraph.space_after or 0) / 12700
        if len(elements) > 500:
            raise ValueError("Too much slide content to preview.")
        try:
            background = _color(slide.background.fill.fore_color, "#FFFFFF")
        except (TypeError, AttributeError):
            background = "#FFFFFF"
        notes = (
            slide.notes_slide.notes_text_frame.text
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame
            else ""
        )
        pages.append(
            PresentationSlide(
                number=index + 1,
                title=title[:500],
                notes=notes[:8000],
                background=background,
                elements=elements,
            )
        )
    return PresentationPreviewResponse(
        title=(prs.core_properties.title or pages[0].title)[:500],
        width=float(prs.slide_width) / 12700,
        height=float(prs.slide_height) / 12700,
        slides_count=len(pages),
        slides=pages,
        warnings=sorted(warnings),
    )


async def presentation_preview(key: str) -> PresentationPreviewResponse:
    from app.engines.storage import storage_engine

    async with asyncio.timeout(15), _slots:
        data = await storage_engine.download_bytes(key, max_bytes=MAX_SOURCE_BYTES)
        with tempfile.TemporaryDirectory(prefix="pansgpt-presentation-") as directory:
            source, target = Path(directory) / "source.pptx", Path(directory) / "preview.json"
            source.write_bytes(data)
            process = subprocess.Popen(
                [sys.executable, "-m", "app.engines.presentations", str(source), str(target)],
                cwd=Path(__file__).resolve().parents[2],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
            )
            try:
                await asyncio.to_thread(process.wait, timeout=8)
            finally:
                if process.poll() is None:
                    process.kill()
                    await asyncio.to_thread(process.wait)
            if process.returncode == 2:
                raise ValueError("Presentation preview is unavailable.")
            if process.returncode or not target.exists():
                raise RuntimeError("The presentation preview worker failed.")
            if target.stat().st_size > 2 * 1024 * 1024:
                raise ValueError("The presentation preview is too large.")
            return PresentationPreviewResponse.model_validate_json(target.read_bytes())


if __name__ == "__main__":
    try:
        result = read_presentation(Path(sys.argv[1]).read_bytes())
    except (ValueError, KeyError, zipfile.BadZipFile, XMLSyntaxError):
        sys.exit(2)
    Path(sys.argv[2]).write_text(result.model_dump_json(), encoding="utf-8")
