# ==============================================================================  # [OFFICE PDF FIX]
# Verification Suite for LibreOffice Office-to-PDF Converter & Worker Pipeline  # [OFFICE PDF FIX]
# ==============================================================================  # [OFFICE PDF FIX]

import asyncio  # [OFFICE PDF FIX]
import io  # [OFFICE PDF FIX]
import os  # [OFFICE PDF FIX]
import unittest.mock  # [OFFICE PDF FIX]
from pathlib import Path  # [OFFICE PDF FIX]

import fitz  # [OFFICE PDF FIX]
import pytest  # [OFFICE PDF FIX]
from arq import Retry  # [OFFICE PDF FIX]

from app.engines.converter import (  # [OFFICE PDF FIX]
    ConversionError,  # [OFFICE PDF FIX]
    convert_office_to_pdf,  # [OFFICE PDF FIX]
)  # [OFFICE PDF FIX]
from workers.tasks import ingest_document_job  # [OFFICE PDF FIX]


def _create_minimal_pdf_bytes(page_count: int = 1) -> bytes:  # [OFFICE PDF FIX]
    doc = fitz.open()  # [OFFICE PDF FIX]
    for _ in range(page_count):  # [OFFICE PDF FIX]
        page = doc.new_page(width=595, height=842)  # [OFFICE PDF FIX]
        page.insert_text((50, 50), "Test PDF Content", fontsize=12)  # [OFFICE PDF FIX]
    data = doc.tobytes()  # [OFFICE PDF FIX]
    doc.close()  # [OFFICE PDF FIX]
    return data  # [OFFICE PDF FIX]


# ------------------------------------------------------------------------------  # [OFFICE PDF FIX]
# 1. Converter Unit Tests (Mocked Subprocess, Fast, No Live Services)  # [OFFICE PDF FIX]
# ------------------------------------------------------------------------------  # [OFFICE PDF FIX]
@pytest.mark.asyncio  # [OFFICE PDF FIX]
async def test_converter_returns_bytes_on_success(monkeypatch):  # [OFFICE PDF FIX]
    """(a) Test converter returns valid PDF bytes on subprocess success."""  # [OFFICE PDF FIX]
    monkeypatch.setattr(
        "app.engines.converter.find_soffice_binary", lambda: "/usr/bin/soffice"
    )  # [OFFICE PDF FIX]

    expected_pdf = _create_minimal_pdf_bytes(1)  # [OFFICE PDF FIX]

    async def mock_subprocess_exec(*cmd, **kwargs):  # [OFFICE PDF FIX]
        outdir = Path(cmd[cmd.index("--outdir") + 1])  # [OFFICE PDF FIX]
        out_pdf = outdir / "input.pdf"  # [OFFICE PDF FIX]
        out_pdf.write_bytes(expected_pdf)  # [OFFICE PDF FIX]

        mock_proc = unittest.mock.AsyncMock()  # [OFFICE PDF FIX]
        mock_proc.returncode = 0  # [OFFICE PDF FIX]
        mock_proc.communicate.return_value = (b"", b"")  # [OFFICE PDF FIX]
        return mock_proc  # [OFFICE PDF FIX]

    monkeypatch.setattr(asyncio, "create_subprocess_exec", mock_subprocess_exec)  # [OFFICE PDF FIX]

    result = await convert_office_to_pdf(b"dummy office content", "docx")  # [OFFICE PDF FIX]
    assert result == expected_pdf  # [OFFICE PDF FIX]


@pytest.mark.asyncio  # [OFFICE PDF FIX]
async def test_converter_soffice_missing_raises_conversion_error(monkeypatch):  # [OFFICE PDF FIX]
    """(b) Test converter raises ConversionError when soffice executable is not found."""  # [OFFICE PDF FIX]
    monkeypatch.setattr(
        "app.engines.converter.find_soffice_binary", lambda: None
    )  # [OFFICE PDF FIX]

    with pytest.raises(
        ConversionError, match="LibreOffice soffice executable not found"
    ):  # [OFFICE PDF FIX]
        await convert_office_to_pdf(b"dummy office bytes", "pptx")  # [OFFICE PDF FIX]


@pytest.mark.asyncio  # [OFFICE PDF FIX]
async def test_converter_timeout_kills_process_and_raises_conversion_error(
    monkeypatch,
):  # [OFFICE PDF FIX]
    """(c) Test converter kills subprocess and raises ConversionError upon timeout."""  # [OFFICE PDF FIX]
    monkeypatch.setattr(
        "app.engines.converter.find_soffice_binary", lambda: "/usr/bin/soffice"
    )  # [OFFICE PDF FIX]

    mock_proc = unittest.mock.AsyncMock()  # [OFFICE PDF FIX]
    mock_proc.kill = unittest.mock.MagicMock()  # [OFFICE PDF FIX]
    mock_proc.wait = unittest.mock.AsyncMock()  # [OFFICE PDF FIX]

    async def hanging_communicate():  # [OFFICE PDF FIX]
        await asyncio.sleep(10.0)  # [OFFICE PDF FIX]
        return (b"", b"")  # [OFFICE PDF FIX]

    mock_proc.communicate = hanging_communicate  # [OFFICE PDF FIX]

    async def mock_subprocess_exec(*cmd, **kwargs):  # [OFFICE PDF FIX]
        return mock_proc  # [OFFICE PDF FIX]

    monkeypatch.setattr(asyncio, "create_subprocess_exec", mock_subprocess_exec)  # [OFFICE PDF FIX]

    # Monkeypatch wait_for with a fast timeout simulation  # [OFFICE PDF FIX]
    orig_wait_for = asyncio.wait_for  # [OFFICE PDF FIX]

    async def fast_wait_for(fut, timeout):  # [OFFICE PDF FIX]
        return await orig_wait_for(fut, timeout=0.01)  # [OFFICE PDF FIX]

    monkeypatch.setattr(asyncio, "wait_for", fast_wait_for)  # [OFFICE PDF FIX]

    with pytest.raises(ConversionError, match="timed out after 180 seconds"):  # [OFFICE PDF FIX]
        await convert_office_to_pdf(b"dummy docx bytes", "docx")  # [OFFICE PDF FIX]

    assert mock_proc.kill.called  # [OFFICE PDF FIX]


# ------------------------------------------------------------------------------  # [OFFICE PDF FIX]
# 2. Worker Office Conversion Flow Tests  # [OFFICE PDF FIX]
# ------------------------------------------------------------------------------  # [OFFICE PDF FIX]
@pytest.mark.asyncio  # [OFFICE PDF FIX]
async def test_worker_converts_office_doc_and_ingests_converted_key(
    monkeypatch,
):  # [OFFICE PDF FIX]
    """(d) The worker converts an Office document, stores converted_key, and ingests the converted key."""  # [OFFICE PDF FIX]
    doc_id = "0851a9c3-a5d9-4969-870d-08849802f93c"  # [OFFICE PDF FIX]
    orig_key = "universities/uni1/courses/pcl/original/lecture.docx"  # [OFFICE PDF FIX]
    converted_pdf_bytes = _create_minimal_pdf_bytes(page_count=3)  # [OFFICE PDF FIX]

    monkeypatch.setattr(
        "workers.tasks._claim_document", unittest.mock.AsyncMock(return_value=True)
    )  # [OFFICE PDF FIX]
    monkeypatch.setattr(
        "workers.tasks._update_document_status", unittest.mock.AsyncMock()
    )  # [OFFICE PDF FIX]

    mock_meta = {  # [OFFICE PDF FIX]
        "id": doc_id,  # [OFFICE PDF FIX]
        "storage_key": orig_key,  # [OFFICE PDF FIX]
        "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",  # [OFFICE PDF FIX]
        "converted_key": None,  # [OFFICE PDF FIX]
        "university_id": "uni1",  # [OFFICE PDF FIX]
        "course_code": "pcl",  # [OFFICE PDF FIX]
        "page_count": None,  # [OFFICE PDF FIX]
    }  # [OFFICE PDF FIX]
    monkeypatch.setattr(
        "workers.tasks._get_document_meta", unittest.mock.AsyncMock(return_value=mock_meta)
    )  # [OFFICE PDF FIX]

    updated_converted_info = []  # [OFFICE PDF FIX]

    async def mock_update_converted(doc_id_arg, key_arg, page_count=None):  # [OFFICE PDF FIX]
        updated_converted_info.append((doc_id_arg, key_arg, page_count))  # [OFFICE PDF FIX]

    monkeypatch.setattr(
        "workers.tasks._update_document_converted_info", mock_update_converted
    )  # [OFFICE PDF FIX]

    from app.engines.storage import storage_engine  # [OFFICE PDF FIX]

    monkeypatch.setattr(
        storage_engine, "download_bytes", unittest.mock.AsyncMock(return_value=b"fake docx")
    )  # [OFFICE PDF FIX]
    uploaded_files = []  # [OFFICE PDF FIX]

    async def mock_upload(key, data, content_type="application/pdf"):  # [OFFICE PDF FIX]
        uploaded_files.append((key, data, content_type))  # [OFFICE PDF FIX]

    monkeypatch.setattr(storage_engine, "upload_bytes", mock_upload)  # [OFFICE PDF FIX]

    monkeypatch.setattr(  # [OFFICE PDF FIX]
        "workers.tasks.convert_office_to_pdf",  # [OFFICE PDF FIX]
        unittest.mock.AsyncMock(return_value=converted_pdf_bytes),  # [OFFICE PDF FIX]
    )  # [OFFICE PDF FIX]

    ingested_keys = []  # [OFFICE PDF FIX]

    from app.engines.ingestion import document_ingestion_engine  # [OFFICE PDF FIX]

    async def mock_ingest_r2(document_id, storage_key):  # [OFFICE PDF FIX]
        ingested_keys.append((document_id, storage_key))  # [OFFICE PDF FIX]
        return True  # [OFFICE PDF FIX]

    monkeypatch.setattr(
        document_ingestion_engine, "ingest_document_from_r2", mock_ingest_r2
    )  # [OFFICE PDF FIX]

    success = await ingest_document_job({}, doc_id, orig_key)  # [OFFICE PDF FIX]
    assert success is True  # [OFFICE PDF FIX]

    expected_converted_key = (
        f"universities/uni1/courses/pcl/converted/{doc_id}.pdf"  # [OFFICE PDF FIX]
    )
    assert len(uploaded_files) == 1  # [OFFICE PDF FIX]
    assert uploaded_files[0][0] == expected_converted_key  # [OFFICE PDF FIX]
    assert uploaded_files[0][1] == converted_pdf_bytes  # [OFFICE PDF FIX]

    assert len(updated_converted_info) == 1  # [OFFICE PDF FIX]
    assert updated_converted_info[0] == (doc_id, expected_converted_key, 3)  # [OFFICE PDF FIX]

    assert len(ingested_keys) == 1  # [OFFICE PDF FIX]
    assert ingested_keys[0] == (doc_id, expected_converted_key)  # [OFFICE PDF FIX]
    assert ingested_keys[0][1] != orig_key  # [OFFICE PDF FIX]


@pytest.mark.asyncio  # [OFFICE PDF FIX]
async def test_worker_conversion_error_does_not_fallback_to_original(
    monkeypatch,
):  # [OFFICE PDF FIX]
    """(e) ConversionError fails or retries the job and NEVER falls back to ingesting the original Office file."""  # [OFFICE PDF FIX]
    doc_id = "14778083-f526-4fc9-9116-ca9b37ed5944"  # [OFFICE PDF FIX]
    orig_key = "universities/uni1/courses/pcl/original/slides.pptx"  # [OFFICE PDF FIX]

    monkeypatch.setattr(
        "workers.tasks._claim_document", unittest.mock.AsyncMock(return_value=True)
    )  # [OFFICE PDF FIX]
    monkeypatch.setattr(
        "workers.tasks._update_document_status", unittest.mock.AsyncMock()
    )  # [OFFICE PDF FIX]

    mock_meta = {  # [OFFICE PDF FIX]
        "id": doc_id,  # [OFFICE PDF FIX]
        "storage_key": orig_key,  # [OFFICE PDF FIX]
        "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",  # [OFFICE PDF FIX]
        "converted_key": None,  # [OFFICE PDF FIX]
        "university_id": "uni1",  # [OFFICE PDF FIX]
        "course_code": "pcl",  # [OFFICE PDF FIX]
        "page_count": None,  # [OFFICE PDF FIX]
    }  # [OFFICE PDF FIX]
    monkeypatch.setattr(
        "workers.tasks._get_document_meta", unittest.mock.AsyncMock(return_value=mock_meta)
    )  # [OFFICE PDF FIX]

    from app.engines.storage import storage_engine  # [OFFICE PDF FIX]

    monkeypatch.setattr(
        storage_engine, "download_bytes", unittest.mock.AsyncMock(return_value=b"fake pptx")
    )  # [OFFICE PDF FIX]

    monkeypatch.setattr(  # [OFFICE PDF FIX]
        "workers.tasks.convert_office_to_pdf",  # [OFFICE PDF FIX]
        unittest.mock.AsyncMock(
            side_effect=ConversionError("LibreOffice crashed")
        ),  # [OFFICE PDF FIX]
    )  # [OFFICE PDF FIX]

    ingested_calls = []  # [OFFICE PDF FIX]

    from app.engines.ingestion import document_ingestion_engine  # [OFFICE PDF FIX]

    monkeypatch.setattr(  # [OFFICE PDF FIX]
        document_ingestion_engine,  # [OFFICE PDF FIX]
        "ingest_document_from_r2",  # [OFFICE PDF FIX]
        unittest.mock.AsyncMock(
            side_effect=lambda *a, **kw: ingested_calls.append(a)
        ),  # [OFFICE PDF FIX]
    )  # [OFFICE PDF FIX]

    with pytest.raises(Retry):  # [OFFICE PDF FIX]
        await ingest_document_job({"job_try": 1}, doc_id, orig_key)  # [OFFICE PDF FIX]

    assert len(ingested_calls) == 0, (
        "Ingestion must NOT run on original file when conversion fails!"
    )  # [OFFICE PDF FIX]


@pytest.mark.asyncio  # [OFFICE PDF FIX]
async def test_worker_plain_pdf_ingested_unchanged(monkeypatch):  # [OFFICE PDF FIX]
    """(f) A plain PDF skips Office conversion and is ingested unchanged."""  # [OFFICE PDF FIX]
    doc_id = "e7a6a964-d4d2-4a97-a721-72494cbd08a7"  # [OFFICE PDF FIX]
    orig_key = "universities/uni1/courses/pcl/original/handout.pdf"  # [OFFICE PDF FIX]

    monkeypatch.setattr(
        "workers.tasks._claim_document", unittest.mock.AsyncMock(return_value=True)
    )  # [OFFICE PDF FIX]
    monkeypatch.setattr(
        "workers.tasks._update_document_status", unittest.mock.AsyncMock()
    )  # [OFFICE PDF FIX]

    mock_meta = {  # [OFFICE PDF FIX]
        "id": doc_id,  # [OFFICE PDF FIX]
        "storage_key": orig_key,  # [OFFICE PDF FIX]
        "mime_type": "application/pdf",  # [OFFICE PDF FIX]
        "converted_key": None,  # [OFFICE PDF FIX]
        "university_id": "uni1",  # [OFFICE PDF FIX]
        "course_code": "pcl",  # [OFFICE PDF FIX]
        "page_count": 5,  # [OFFICE PDF FIX]
    }  # [OFFICE PDF FIX]
    monkeypatch.setattr(
        "workers.tasks._get_document_meta", unittest.mock.AsyncMock(return_value=mock_meta)
    )  # [OFFICE PDF FIX]

    convert_called = []  # [OFFICE PDF FIX]
    monkeypatch.setattr(  # [OFFICE PDF FIX]
        "workers.tasks.convert_office_to_pdf",  # [OFFICE PDF FIX]
        unittest.mock.AsyncMock(
            side_effect=lambda *a, **kw: convert_called.append(a)
        ),  # [OFFICE PDF FIX]
    )  # [OFFICE PDF FIX]

    ingested_keys = []  # [OFFICE PDF FIX]

    from app.engines.ingestion import document_ingestion_engine  # [OFFICE PDF FIX]

    async def mock_ingest(document_id, storage_key):  # [OFFICE PDF FIX]
        ingested_keys.append((document_id, storage_key))  # [OFFICE PDF FIX]
        return True  # [OFFICE PDF FIX]

    monkeypatch.setattr(
        document_ingestion_engine, "ingest_document_from_r2", mock_ingest
    )  # [OFFICE PDF FIX]

    success = await ingest_document_job({}, doc_id, orig_key)  # [OFFICE PDF FIX]
    assert success is True  # [OFFICE PDF FIX]
    assert len(convert_called) == 0, (
        "Plain PDF must not invoke Office converter!"
    )  # [OFFICE PDF FIX]
    assert len(ingested_keys) == 1  # [OFFICE PDF FIX]
    assert ingested_keys[0] == (doc_id, orig_key)  # [OFFICE PDF FIX]


# ------------------------------------------------------------------------------  # [OFFICE PDF FIX]
# 3. Live Test (Skipped unless RUN_LIVE_CONVERT_TEST=1)  # [OFFICE PDF FIX]
# ------------------------------------------------------------------------------  # [OFFICE PDF FIX]
@pytest.mark.live  # [OFFICE PDF FIX]
@pytest.mark.skipif(  # [OFFICE PDF FIX]
    os.getenv("RUN_LIVE_CONVERT_TEST") != "1",  # [OFFICE PDF FIX]
    reason="Live conversion test requires RUN_LIVE_CONVERT_TEST=1 and LibreOffice installed.",  # [OFFICE PDF FIX]
)  # [OFFICE PDF FIX]
@pytest.mark.asyncio  # [OFFICE PDF FIX]
async def test_live_office_to_pdf_conversion():  # [OFFICE PDF FIX]
    """Live conversion of synthetic docx and pptx via local LibreOffice."""  # [OFFICE PDF FIX]
    import docx  # [OFFICE PDF FIX]
    import pptx  # [OFFICE PDF FIX]

    # 1. Tiny docx  # [OFFICE PDF FIX]
    doc = docx.Document()  # [OFFICE PDF FIX]
    doc.add_heading("Live Conversion Test Document", 0)  # [OFFICE PDF FIX]
    doc.add_paragraph("First paragraph testing LibreOffice conversion.")  # [OFFICE PDF FIX]
    docx_io = io.BytesIO()  # [OFFICE PDF FIX]
    doc.save(docx_io)  # [OFFICE PDF FIX]
    docx_bytes = docx_io.getvalue()  # [OFFICE PDF FIX]

    pdf_docx = await convert_office_to_pdf(docx_bytes, "docx")  # [OFFICE PDF FIX]
    assert len(pdf_docx) > 0  # [OFFICE PDF FIX]
    pdf_doc = fitz.open(stream=pdf_docx, filetype="pdf")  # [OFFICE PDF FIX]
    assert pdf_doc.page_count >= 1  # [OFFICE PDF FIX]
    pdf_doc.close()  # [OFFICE PDF FIX]

    # 2. Tiny pptx (2 slides)  # [OFFICE PDF FIX]
    prs = pptx.Presentation()  # [OFFICE PDF FIX]
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])  # [OFFICE PDF FIX]
    slide1.shapes.title.text = "Slide 1 Title"  # [OFFICE PDF FIX]
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])  # [OFFICE PDF FIX]
    slide2.shapes.title.text = "Slide 2 Title"  # [OFFICE PDF FIX]
    pptx_io = io.BytesIO()  # [OFFICE PDF FIX]
    prs.save(pptx_io)  # [OFFICE PDF FIX]
    pptx_bytes = pptx_io.getvalue()  # [OFFICE PDF FIX]

    pdf_pptx = await convert_office_to_pdf(pptx_bytes, "pptx")  # [OFFICE PDF FIX]
    assert len(pdf_pptx) > 0  # [OFFICE PDF FIX]
    pdf_prs = fitz.open(stream=pdf_pptx, filetype="pdf")  # [OFFICE PDF FIX]
    assert pdf_prs.page_count == 2  # [OFFICE PDF FIX]
    pdf_prs.close()  # [OFFICE PDF FIX]
