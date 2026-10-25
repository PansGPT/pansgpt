"""Actual-file preview, exact slide totals, isolated parsing and owner boundaries."""

import io
import subprocess
import zipfile
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.engines import presentations
from app.engines.llm import MultiTierLlmEngine
from app.engines.skills import SKILL_TOOL_DEFINITIONS, skill_engine
from app.engines.storage import storage_engine
from app.main import app
from app.models.presentation import PresentationContentSlide, PresentationPlan, PresentationSource
from tests.test_artifact_security import OTHER, OWNER, UNI


@pytest.fixture
def deck():
    slides = [
        PresentationContentSlide(
            title=f"Synthetic topic {i}",
            bullet_points=[f"Learning point {i}", "A second concise point"],
            notes=f"Speaker note {i}",
        )
        for i in range(1, 5)
    ]
    sources = [
        PresentationSource(title="Synthetic reference", url="https://example.test/reference")
    ]
    data, metadata = presentations.build_presentation(
        "Synthetic presentation", slides, "A clear study outline", 5, sources
    )
    return data, metadata, slides


def test_saved_file_has_one_cover_five_slides_notes_and_readable_geometry(deck):
    data, metadata, _ = deck
    file = Presentation(io.BytesIO(data))
    preview = presentations.read_presentation(data)
    assert len(file.slides) == preview.slides_count == metadata["slides_count"] == 5
    assert preview.width == pytest.approx(960, abs=0.01) and preview.height == 540
    assert preview.slides[0].title == "Synthetic presentation"
    assert preview.slides[1].title == "Synthetic topic 1"
    assert "Speaker note 1" in preview.slides[1].notes
    assert "https://example.test/reference" in preview.slides[1].notes
    assert not preview.warnings
    assert any(
        element.kind == "text" and "Learning point 1" in element.text
        for element in preview.slides[1].elements
    )
    assert len(set(page.title for page in preview.slides)) == 5
    for slide in file.slides:
        text = [shape.text for shape in slide.shapes if shape.has_text_frame]
        assert "PANSGPT  /  STUDY PRESENTATION" in text
        assert "PansGPT" in text


def test_preview_reads_current_file_not_a_cached_outline(deck):
    data, _, _ = deck
    file = Presentation(io.BytesIO(data))
    title = next(shape for shape in file.slides[1].shapes if shape.name == "pansgpt-title")
    title.text = "Changed in the actual file"
    saved = io.BytesIO()
    file.save(saved)
    assert presentations.read_presentation(saved.getvalue()).slides[1].title == title.text


def test_count_and_duplicate_cover_errors_do_not_silently_pad_or_truncate(deck):
    _, _, slides = deck
    with pytest.raises(ValueError, match="including one cover"):
        presentations.build_presentation("Synthetic presentation", slides, None, 6, [])
    slides[0].title = "Synthetic presentation"
    with pytest.raises(ValueError, match="cover is added once"):
        presentations.build_presentation("Synthetic presentation", slides, None, 5, [])


async def test_router_count_cannot_be_overridden_by_tool_arguments(deck, monkeypatch):
    _, _, slides = deck
    upload = AsyncMock()
    monkeypatch.setattr(skill_engine, "_safe_r2_upload", upload)
    with pytest.raises(HTTPException) as error:
        await skill_engine.dispatch_skill(
            "create_pptx",
            {
                "title": "Synthetic presentation",
                "slide_count": 6,
                "slides": [slide.model_dump() for slide in slides],
            },
            user_id=OWNER,
            university_id=UNI,
            requested_slide_count=5,
        )
    assert error.value.status_code == 422
    upload.assert_not_awaited()


def test_media_is_reported_without_fetching_or_inventing_it(deck):
    data, _, _ = deck
    file = Presentation(io.BytesIO(data))
    image = io.BytesIO()
    Image.new("RGB", (2, 2), "green").save(image, "PNG")
    image.seek(0)
    file.slides[1].shapes.add_picture(image, Inches(1), Inches(1))
    saved = io.BytesIO()
    file.save(saved)
    preview = presentations.read_presentation(saved.getvalue())
    assert any("media" in warning for warning in preview.warnings)
    assert {item.kind for item in preview.slides[1].elements} <= {"rect", "ellipse", "text"}


@pytest.mark.parametrize("mode", ["parts", "xml", "size"])
def test_unsafe_or_oversized_packages_are_rejected_before_parsing(mode):
    if mode == "size":
        with pytest.raises(ValueError, match="large"):
            presentations.read_presentation(b"x" * (presentations.MAX_SOURCE_BYTES + 1))
        return
    package = io.BytesIO()
    with zipfile.ZipFile(package, "w") as archive:
        if mode == "parts":
            for index in range(601):
                archive.writestr(f"part{index}", "x")
        else:
            archive.writestr("ppt/test.xml", '<!DOCTYPE test [<!ENTITY bad "example">]><test/>')
    with pytest.raises(ValueError, match="limits|Unsafe"):
        presentations.read_presentation(package.getvalue())


@pytest.mark.parametrize(
    "key",
    [
        f"artifacts/{UNI}/{OTHER}/pptx/file.pptx",
        f"artifacts/{OTHER}/{OWNER}/pptx/file.pptx",
        f"artifacts/{UNI}/{OWNER}/pptx/../file.pptx",
        f"artifacts/{UNI}/{OWNER}/pdf/file.pdf",
    ],
)
def test_preview_denies_foreign_and_invalid_keys_before_storage(
    authenticated_student, monkeypatch, key
):
    read = AsyncMock()
    monkeypatch.setattr(storage_engine, "download_bytes", read)
    response = TestClient(app).get("/api/v1/library/artifacts/presentation", params={"key": key})
    assert response.status_code == 404
    read.assert_not_awaited()


def test_owned_preview_uses_isolated_parser_and_private_response(
    authenticated_student, monkeypatch, deck
):
    read = AsyncMock(return_value=deck[0])
    monkeypatch.setattr(storage_engine, "download_bytes", read)
    response = TestClient(app).get(
        "/api/v1/library/artifacts/presentation",
        params={"key": f"artifacts/{UNI}/{OWNER}/pptx/file.pptx"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["slides_count"] == 5
    assert response.json()["slides"][1]["title"] == "Synthetic topic 1"
    assert response.headers["cache-control"] == "private, no-store"
    assert read.call_args.kwargs["max_bytes"] == presentations.MAX_SOURCE_BYTES


async def test_timed_out_parser_is_killed_and_temporary_files_are_removed(monkeypatch, deck):
    monkeypatch.setattr(storage_engine, "download_bytes", AsyncMock(return_value=deck[0]))
    children = []

    class Stalled:
        def __init__(self, args, **kwargs):
            self.path = presentations.Path(args[-2]).parent
            self.killed = False
            children.append(self)

        def wait(self, timeout=None):
            if not self.killed:
                raise subprocess.TimeoutExpired("synthetic-parser", timeout)
            return 1

        def poll(self):
            return 1 if self.killed else None

        def kill(self):
            self.killed = True

    monkeypatch.setattr(presentations.subprocess, "Popen", Stalled)
    with pytest.raises(subprocess.TimeoutExpired):
        await presentations.presentation_preview("synthetic")
    assert children[0].killed and not children[0].path.exists()


def test_corrupt_owned_file_returns_a_preview_error_with_download_still_independent(
    authenticated_student, monkeypatch
):
    monkeypatch.setattr(
        storage_engine, "download_bytes", AsyncMock(return_value=b"not a PowerPoint package")
    )
    response = TestClient(app).get(
        "/api/v1/library/artifacts/presentation",
        params={"key": f"artifacts/{UNI}/{OWNER}/pptx/file.pptx"},
    )
    assert response.status_code == 422
    assert "Download" in response.json()["detail"]


@pytest.mark.parametrize("requested,declared,valid", [(5, 5, True), (5, 6, False), (None, 5, True)])
async def test_document_context_requires_a_semantic_plan_before_upload(
    monkeypatch, deck, requested, declared, valid
):
    engine = MultiTierLlmEngine()
    engine.gemini_key, engine.groq_key, engine.openrouter_key = "fixture-provider", "", ""
    monkeypatch.setattr(engine, "_record_telemetry", AsyncMock())
    monkeypatch.setattr(
        "app.engines.skill_registry.available_tools", AsyncMock(return_value=SKILL_TOOL_DEFINITIONS)
    )
    planner = AsyncMock(return_value=PresentationPlan(slide_count=requested))
    monkeypatch.setattr("app.engines.presentation_planning.resolve_presentation_plan", planner)
    key = f"artifacts/{UNI}/{OWNER}/pptx/synthetic.pptx"
    upload = AsyncMock(return_value=(key, "/api/v1/library/artifacts/download?key=" + key))
    monkeypatch.setattr(skill_engine, "_safe_r2_upload", upload)
    calls = [
        {
            "id": "synthetic-pptx",
            "name": "create_pptx",
            "arguments": {
                "title": "Synthetic presentation",
                "slide_count": declared,
                "slides": [slide.model_dump() for slide in deck[2]],
            },
        }
    ]
    monkeypatch.setattr(engine, "_call_gemma_turn", AsyncMock(return_value=(calls, None)))
    events = [
        event
        async for event in engine.stream_agentic_chat(
            "Study tutor",
            "Turn this document into five slides",
            user_id=OWNER,
            university_id=UNI,
            validate_presentation_request=True,
        )
    ]
    planner.assert_awaited_once()
    if valid:
        upload.assert_awaited_once()
        assert len(Presentation(io.BytesIO(upload.call_args.args[1])).slides) == 5
        assert any(kind == "artifact_ready" for kind, _, _ in events)
    else:
        upload.assert_not_awaited()
        assert not any(kind == "artifact_ready" for kind, _, _ in events)
        assert any(
            kind == "error" and "requested 5" in payload["error"] for kind, payload, _ in events
        )
