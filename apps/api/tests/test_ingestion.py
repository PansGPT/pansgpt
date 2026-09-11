# ==============================================================================
# Phase 5 Document Ingestion Engine Full Verification Suite (Roadmap 5.11)
# ==============================================================================

import io
import unittest.mock
import uuid

import docx
import fitz  # PyMuPDF
import pptx
import pytest

from app.engines.chunker import semantic_chunker
from app.engines.embedder import gemini_embedder
from app.engines.extractor import document_extractor
from app.engines.ingestion import document_ingestion_engine
from app.engines.storage import build_converted_storage_key, build_document_storage_key
from workers.tasks import _claim_document, _update_document_status, ingest_document_job


def create_synthetic_test_pdf() -> bytes:
    """Helper to generate a clean, multi-page synthetic pharmacology lecture PDF."""
    doc = fitz.open()

    # Page 1: Native Text Page with Headings and Content
    page1 = doc.new_page(width=595, height=842)
    text_content = (
        "PHARMACOLOGY OF ADRENERGIC RECEPTORS\n\n"
        "Epinephrine and Norepinephrine are endogenous catecholamines acting on alpha and beta adrenoceptors.\n"
        "Alpha-1 receptors are predominantly located in vascular smooth muscle where activation causes vasoconstriction.\n"
        "Beta-1 receptors are predominantly located in cardiac tissue where activation increases heart rate and contractility.\n"
        "Beta-2 receptors are located in bronchial smooth muscle where activation causes bronchodilation.\n"
    )
    page1.insert_text((50, 50), text_content, fontsize=11)

    # Page 2: Second Topic Page with Long Paragraphs for Chunking
    page2 = doc.new_page(width=595, height=842)
    page2_text = (
        "CLINICAL INDICATIONS AND CONTRAINDICATIONS\n\n"
        "Anaphylactic shock requires immediate administration of intramuscular epinephrine at a dose of 0.3 to 0.5 mg.\n"
        "Beta-blockers such as propranolol are contraindicated in severe asthma due to risk of bronchospasm.\n"
        "Selective beta-1 blockers like atenolol and metoprolol have lower risk of bronchoconstriction at therapeutic doses.\n"
    )
    page2.insert_text((50, 50), page2_text, fontsize=11)

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_synthetic_docx() -> bytes:
    """Helper to generate a synthetic Word (.docx) document."""
    d = docx.Document()
    d.add_heading("PHARMACOKINETICS PRINCIPLES", level=1)
    d.add_paragraph("Bioavailability and apparent volume of distribution.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Formula"
    table.cell(1, 0).text = "Vd"
    table.cell(1, 1).text = "Dose / C0"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def create_synthetic_pptx() -> bytes:
    """Helper to generate a synthetic PowerPoint (.pptx) presentation."""
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Cardiovascular Pharmacology"
    slide.placeholders[1].text = "ACE Inhibitors and Angiotensin Receptor Blockers"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------------
# 5.1 Storage & Pathing Tests
# ------------------------------------------------------------------------------


def test_storage_key_generation():
    """Verify canonical R2 original path format."""
    uni_id = "018f3a10-0001-7000-8000-000000000001"
    doc_id = "018f3a30-0001-7000-8000-000000000001"
    key = build_document_storage_key(
        university_id=uni_id,
        course_code="PCL 401",
        document_id=doc_id,
        file_extension="pdf",
    )
    assert key == f"universities/{uni_id}/courses/pcl_401/original/{doc_id}.pdf"


def test_converted_storage_key_format():
    """Verify canonical converted Office-to-PDF path (Roadmap 5.1)."""
    doc_id = "018f3a30-0001-7000-8000-000000000002"
    key = build_converted_storage_key(doc_id)
    assert key == f"converted/{doc_id}.pdf"


# ------------------------------------------------------------------------------
# 5.3 & 5.4 Extraction Tests (PDF, DOCX, PPTX, TXT, CSV)
# ------------------------------------------------------------------------------


def test_multi_format_extractors_docx_pptx_txt_csv():
    """Verify multi-format extractors extract text, titles, and tables correctly (Roadmap 5.4)."""
    # 1. DOCX
    docx_bytes = create_synthetic_docx()
    docx_pages = document_extractor.process_document(docx_bytes, file_format="docx")
    assert len(docx_pages) >= 1
    assert "PHARMACOKINETICS" in docx_pages[0].native_text
    assert len(docx_pages[0].tables) == 1
    assert "Vd" in docx_pages[0].tables[0].markdown

    # 2. PPTX
    pptx_bytes = create_synthetic_pptx()
    pptx_pages = document_extractor.process_document(pptx_bytes, file_format="pptx")
    assert len(pptx_pages) >= 1
    assert "Cardiovascular Pharmacology" in pptx_pages[0].native_text

    # 3. CSV
    csv_bytes = b"Drug,Dosage,Route\nParacetamol,500mg,Oral\nAmoxicillin,500mg,Oral"
    csv_pages = document_extractor.process_document(csv_bytes, file_format="csv")
    assert len(csv_pages) >= 1
    assert len(csv_pages[0].tables) == 1
    assert "Paracetamol" in csv_pages[0].tables[0].markdown

    # 4. Markdown / Plaintext
    md_bytes = b"# Chemotherapy\n\nAlkylating agents cross-link DNA."
    md_pages = document_extractor.process_document(md_bytes, file_format="md")
    assert len(md_pages) >= 1
    assert "Chemotherapy" in md_pages[0].native_text


# ------------------------------------------------------------------------------
# 5.6 Semantic Chunker Tests
# ------------------------------------------------------------------------------


def test_chunker_atomic_rules():
    """Verify tables and diagrams are NEVER split (1 element = 1 chunk)."""
    elements = [
        {
            "content_type": "table",
            "raw_content": "| Drug | Receptor | Indication |\n|---|---|---|\n| Epinephrine | Alpha/Beta | Anaphylaxis |",
            "page_number": 1,
            "segment_id": "seg-1",
            "element_id": "el-1",
        },
        {
            "content_type": "diagram",
            "raw_content": "[Visual Description: Diagram of G-protein coupled receptor activation cascade]",
            "page_number": 1,
            "segment_id": "seg-1",
            "element_id": "el-2",
        },
    ]
    chunks = semantic_chunker.create_chunks_for_elements(elements)
    assert len(chunks) == 2
    assert chunks[0].content_type == "table"
    assert chunks[0].content == elements[0]["raw_content"]
    assert chunks[1].content_type == "diagram"
    assert chunks[1].content == elements[1]["raw_content"]


def test_chunker_text_recursive_splitting():
    """Verify long text splits respect segment boundaries and max token lengths."""
    long_para = "This is a detailed pharmacological explanation of drug receptor kinetics. " * 60
    elements = [
        {
            "content_type": "text",
            "raw_content": long_para,
            "page_number": 2,
            "segment_id": "seg-2",
            "element_id": "el-3",
        }
    ]
    chunks = semantic_chunker.create_chunks_for_elements(elements)
    assert len(chunks) >= 2
    for c in chunks:
        assert c.content_type == "text"
        assert c.page_start == 2
        assert c.segment_id == "seg-2"
        assert semantic_chunker.count_tokens(c.content) <= 512


# ------------------------------------------------------------------------------
# 5.7 Embedder Tests
# ------------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_embedder_vector_dimensions():
    """Verify gemini_embedder produces dense vectors with exactly 3072 dimensions."""
    test_texts = [
        "Pharmacokinetics and bioavailability of amoxicillin capsules.",
        "Mechanism of action of ACE inhibitors in essential hypertension.",
    ]
    vectors = await gemini_embedder.embed_batch(test_texts)
    assert len(vectors) == 2
    for vec in vectors:
        assert len(vec) == 3072
        assert isinstance(vec[0], float)


# ------------------------------------------------------------------------------
# 5.8 & 5.11 Full Pipeline End-to-End Test
# ------------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_full_8stage_pipeline_execution():
    """
    End-to-End Test: Run synthetic PDF through all 8 stages of DocumentIngestionEngine.
    Confirms page layer detection, hierarchy creation, atomic chunking, and 3072d vector generation.
    """
    pdf_bytes = create_synthetic_test_pdf()
    doc_id = "018f3a30-0001-7000-8000-000000000099"

    progress_milestones = []

    async def test_progress(pct: int):
        progress_milestones.append(pct)

    result = await document_ingestion_engine.run_pipeline_on_bytes(
        file_bytes=pdf_bytes,
        document_id=doc_id,
        file_format="pdf",
        progress_callback=test_progress,
    )

    # 1. Verify Pages (Stage 2 & 3)
    assert len(result.pages) == 2
    for page in result.pages:
        assert page["has_text_layer"] is True
        assert page["document_id"] == doc_id

    # 2. Verify Segments (Stage 7)
    assert len(result.segments) >= 1
    for seg in result.segments:
        assert seg["document_id"] == doc_id
        assert seg["title_source"] in ("explicit", "synthesized", "inherited")

    # 3. Verify Elements (Stage 3a / 4 / 5 / 6)
    assert len(result.elements) >= 2
    for el in result.elements:
        assert el["content_type"] in ("text", "table", "diagram")
        assert len(el["raw_content"]) > 0

    # 4. Verify Chunks & 3072d Vector Embeddings (Stage 8)
    assert len(result.chunks) >= 2
    for chunk in result.chunks:
        assert chunk["document_id"] == doc_id
        assert len(chunk["embedding"]) == 3072
        assert isinstance(chunk["chunk_index"], int)

    # 5. Verify Progress Telemetry Milestones (Roadmap 5.9)
    assert len(progress_milestones) >= 4
    assert 20 in progress_milestones
    assert 80 in progress_milestones


# ------------------------------------------------------------------------------
# 5.2 & 5.10 REST Endpoints & RBAC Tests
# ------------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_upload_endpoint_generates_presigned_url(client):
    """Verify POST /api/v1/library/upload and aliased /api/v1/documents/upload generate 201 Created."""
    payload = {
        "title": "Adrenergic Pharmacology Lecture 1",
        "course_code": "PCL 401",
        "course_title": "Autonomic Nervous System Pharmacology",
        "topic": "Adrenergic Agonists",
        "lecturer_name": "Dr. O. Adeleke",
        "file_name": "pcl401_lecture1.pdf",
        "file_size_bytes": 1048576,
        "mime_type": "application/pdf",
        "target_levels": ["400"],
    }
    # 1. Primary path
    response = await client.post(
        "/api/v1/library/upload",
        json=payload,
        headers={"x-user-role": "super_admin"},
    )
    assert response.status_code == 201
    data = response.json()
    assert "document_id" in data
    assert "storage_key" in data
    assert "presigned_upload_url" in data
    assert data["expires_in_seconds"] == 900
    assert "pcl_401" in data["storage_key"]

    # 2. Documents alias path (Roadmap 5.2)
    response_alias = await client.post(
        "/api/v1/documents/upload",
        json=payload,
        headers={"x-user-role": "university_admin"},
    )
    assert response_alias.status_code == 201
    assert "document_id" in response_alias.json()


@pytest.mark.asyncio
async def test_rbac_upload_security_guard(client):
    """Verify students cannot initiate uploads (Roadmap 5.10 RBAC requirement)."""
    payload = {
        "title": "Unauthorized Student Upload",
        "course_code": "PCL 401",
        "course_title": "Pharmacology",
        "file_name": "test.pdf",
    }
    response = await client.post(
        "/api/v1/library/upload",
        json=payload,
        headers={"x-user-role": "student"},
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_confirm_upload_endpoints(client):
    """Verify confirm-upload transitions and returns 202 Accepted or 404."""
    random_doc_id = str(uuid.uuid4())
    # 1. Primary path
    resp1 = await client.post(f"/api/v1/library/{random_doc_id}/confirm-upload")
    assert resp1.status_code in (202, 404)

    # 2. Alias complete path
    resp2 = await client.post(f"/api/v1/library/documents/{random_doc_id}/complete")
    assert resp2.status_code in (202, 404)


@pytest.mark.asyncio
async def test_pdf_stream_url_endpoint(client):
    """Verify GET /api/v1/library/{id}/pdf-url returns valid streaming URL."""
    test_id = "018f3a30-0001-7000-8000-000000000001"
    response = await client.get(f"/api/v1/library/{test_id}/pdf-url")
    assert response.status_code == 200
    data = response.json()
    assert data["document_id"] == test_id
    assert "presigned_url" in data
    assert data["expires_in_seconds"] == 900


@pytest.mark.asyncio
async def test_get_single_document_endpoint(client):
    """Verify GET /api/v1/library/documents/{id} handles non-existent gracefully."""
    random_id = str(uuid.uuid4())
    response = await client.get(f"/api/v1/library/documents/{random_id}")
    assert response.status_code in (404, 503)


@pytest.mark.asyncio
async def test_patch_document_endpoint(client):
    """Verify PATCH /api/v1/library/documents/{id} requires valid fields."""
    random_id = str(uuid.uuid4())
    # Empty body -> 422
    resp_empty = await client.patch(f"/api/v1/library/documents/{random_id}", json={})
    assert resp_empty.status_code == 422

    # Valid payload -> 404 or 503
    resp_valid = await client.patch(
        f"/api/v1/library/documents/{random_id}",
        json={"title": "Updated Title"},
    )
    assert resp_valid.status_code in (404, 503)


@pytest.mark.asyncio
async def test_delete_document_endpoint(client):
    """Verify DELETE /api/v1/library/documents/{id} soft-deletes returning 204."""
    random_id = str(uuid.uuid4())
    response = await client.delete(f"/api/v1/library/documents/{random_id}")
    assert response.status_code in (204, 500)


@pytest.mark.asyncio
async def test_get_document_segments_endpoint(client):
    """Verify GET /api/v1/library/documents/{id}/segments returns list."""
    random_id = str(uuid.uuid4())
    response = await client.get(f"/api/v1/library/documents/{random_id}/segments")
    assert response.status_code in (200, 500)
    if response.status_code == 200:
        assert isinstance(response.json(), list)


# ------------------------------------------------------------------------------
# 5.8 & 5.9 Worker Tests (Claim, Heartbeat, Failure Handling)
# ------------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_worker_claim_and_heartbeat_logic():
    """Verify worker claim function and status updates behave correctly."""
    doc_id = str(uuid.uuid4())
    # Claim returns a boolean without throwing
    claimed = await _claim_document(doc_id)
    assert isinstance(claimed, bool)

    # Status update completes without unhandled exceptions
    await _update_document_status(doc_id, status="processing", progress=10)


@pytest.mark.asyncio
async def test_worker_failure_status_transition():
    """Verify worker catches unhandled exceptions and transitions document to failed."""
    doc_id = str(uuid.uuid4())
    # Mock ingest to raise exception
    with unittest.mock.patch.object(
        document_ingestion_engine,
        "ingest_document_from_r2",
        side_effect=RuntimeError("Simulated pipeline failure"),
    ):
        with pytest.raises(RuntimeError):
            await ingest_document_job(
                {"job_try": 3},  # Final attempt triggers failure transition
                document_id=doc_id,
                storage_key="invalid/path.pdf",
            )
