"""Generated file bytes, owner authorization and file-only responses; no external services."""

import io
from collections.abc import AsyncIterator
from unittest.mock import AsyncMock
from urllib.parse import quote

import pytest
from botocore.exceptions import ClientError
from fastapi.testclient import TestClient
from pptx import Presentation

from app.core.dependencies import UserContext
from app.engines.llm import MultiTierLlmEngine
from app.engines.skills import SKILL_TOOL_DEFINITIONS, skill_engine
from app.engines.storage import storage_engine
from app.main import app
from app.models.presentation import PresentationSource
from tests.test_artifact_security import OWNER, UNI


@pytest.mark.parametrize("storage_fails", [False, True])
async def test_presentation_is_saved_before_card_and_not_repeated_in_chat(
    authenticated_student: UserContext, monkeypatch: pytest.MonkeyPatch, storage_fails: bool
) -> None:
    engine = MultiTierLlmEngine()
    engine.gemini_key = "fixture-provider"
    engine.groq_key = engine.openrouter_key = ""
    monkeypatch.setattr(engine, "_record_telemetry", AsyncMock())
    monkeypatch.setattr(
        "app.engines.skill_registry.available_tools", AsyncMock(return_value=SKILL_TOOL_DEFINITIONS)
    )
    monkeypatch.setattr(type(storage_engine), "is_configured", property(lambda _: True))
    saved: dict[str, bytes] = {}

    async def upload(key: str, data: bytes, content_type: str) -> None:
        if storage_fails:
            raise RuntimeError("synthetic storage outage")
        saved[key] = data

    monkeypatch.setattr(storage_engine, "upload_bytes", upload)

    async def preamble() -> AsyncIterator[str]:
        yield "Slide 1: Synthetic study outline. "
        yield "Slide 2: This content belongs in the file."

    call = AsyncMock(
        return_value=(
            [
                {
                    "id": "synthetic-call",
                    "name": "create_pptx",
                    "arguments": {
                        "title": "Synthetic presentation",
                        "overview": "The deck introduces the synthetic study topic and its learning focus. It provides a concise outline for revision.",
                        "sources": [{"title": "Invented source", "url": "https://invented.test"}],
                        "slides": [
                            {"title": "Synthetic content", "bullet_points": ["Fixture bullet"]},
                        ],
                    },
                },
            ],
            preamble(),
        )
    )
    monkeypatch.setattr(engine, "_call_gemma_turn", call)
    events = [
        event
        async for event in engine.stream_agentic_chat(
            "Study tutor",
            "Make a presentation",
            user_id=OWNER,
            university_id=UNI,
            presentation_sources=[
                PresentationSource(title="Course monograph"),
                PresentationSource(
                    title="Synthetic reference", url="https://example.test/reference"
                ),
            ],
        )
    ]
    text = "".join(payload["token"] for kind, payload, _ in events if kind == "text_chunk")
    artifacts = [payload for kind, payload, _ in events if kind == "artifact_ready"]
    assert call.await_count == 1
    if storage_fails:
        assert not artifacts and not text and not saved
        assert any(kind == "tool_end" and payload.status == "error" for kind, payload, _ in events)
        assert any(kind == "error" and payload["retryable"] for kind, payload, _ in events)
        return
    assert text.startswith("The deck introduces the synthetic study topic and its learning focus.")
    assert "Your 2-slide presentation is ready" not in text
    assert "The deck introduces the synthetic study topic and its learning focus." in text
    assert "It provides a concise outline for revision." in text
    assert "- Course monograph" in text
    assert "[Synthetic reference](<https://example.test/reference>)" in text
    assert "Invented source" not in text and "invented.test" not in text
    assert "Fixture bullet" not in text and "Slide 1:" not in text
    assert len(artifacts) == 1
    artifact = artifacts[0]
    original = saved[artifact.storage_key]
    assert len(Presentation(io.BytesIO(original)).slides) == 2
    assert (
        "Course monograph"
        in Presentation(io.BytesIO(original)).slides[0].notes_slide.notes_text_frame.text
    )
    assert artifact.content["overview"] in text
    monkeypatch.setattr(storage_engine, "download_bytes", AsyncMock(return_value=original))
    response = TestClient(app).get(artifact.download_url)
    assert response.status_code == 200
    assert response.content == original
    assert response.headers["content-type"].startswith(artifact.mime_type)
    assert response.headers["cache-control"] == "private, no-store"
    assert "attachment" in response.headers["content-disposition"]
    authenticated_student["id"] = "018f3a10-0001-7000-8000-000000000002"
    assert TestClient(app).get(artifact.download_url).status_code == 404


async def test_unconfigured_storage_cannot_advertise_a_download(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(type(storage_engine), "is_configured", property(lambda _: False))
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as error:
        await skill_engine._safe_r2_upload("synthetic", b"data", "text/plain")
    assert error.value.status_code == 503


@pytest.mark.parametrize("missing", [False, True])
def test_missing_file_and_storage_outage_are_distinct(
    authenticated_student: UserContext, monkeypatch: pytest.MonkeyPatch, missing: bool
) -> None:
    error = (
        ClientError({"Error": {"Code": "NoSuchKey"}}, "GetObject")
        if missing
        else RuntimeError("outage")
    )
    monkeypatch.setattr(storage_engine, "download_bytes", AsyncMock(side_effect=error))
    response = TestClient(app).get(
        "/api/v1/library/artifacts/download",
        params={"key": f"artifacts/{UNI}/{OWNER}/pptx/fixture.pptx"},
    )
    assert response.status_code == (404 if missing else 503)


def test_unicode_filename_downloads_without_header_encoding_404(
    authenticated_student: UserContext,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(storage_engine, "download_bytes", AsyncMock(return_value=b"fixture"))
    title = "\u03b2 blockers"
    response = TestClient(app).get(
        "/api/v1/library/artifacts/download",
        params={
            "key": f"artifacts/{UNI}/{OWNER}/pptx/fixture.pptx",
            "filename": title,
        },
    )
    assert response.status_code == 200
    assert (
        f"filename*=UTF-8''{quote(title.replace(' ', '_') + '.pptx')}"
        in response.headers["content-disposition"]
    )
